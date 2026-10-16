import io
import sys
import uuid
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app import create_app  # noqa: E402
from app.config import Config  # noqa: E402
from app.db import connect  # noqa: E402
from migrations.migrate import run as run_migrations  # noqa: E402
from seed.seed import seed  # noqa: E402

TEST_DB = "finance_kg_test"


@pytest.fixture(scope="session")
def cfg(tmp_path_factory):
    c = Config.from_env(DB_NAME=TEST_DB, PROCESSING_MODE="sync", UPLOAD_DIR=str(tmp_path_factory.mktemp("uploads")),
                        MAX_UPLOAD_MB=1, LLM_PROVIDER="none", TESTING=True, SECRET_KEY="test-secret-key-0123456789abcdef-xyz")
    run_migrations(c, reset=True)
    seed(c)
    return c


@pytest.fixture()
def app(cfg):
    conn = connect(cfg)
    with conn.cursor() as cur:  # fresh per-user state for every test (KG0 tables are untouched)
        cur.execute("DELETE FROM users")
        cur.execute("DELETE FROM audit_logs")
    conn.close()
    return create_app(cfg)


@pytest.fixture()
def client(app):
    return app.test_client()


class Api:
    """Small helper that wraps the Flask test client with auth + common workflow steps."""

    def __init__(self, client, email=None):
        self.c = client
        self.email = email or f"u{uuid.uuid4().hex[:8]}@example.com"
        r = client.post("/api/auth/register", json={"email": self.email, "display_name": "Test User", "password": "Passw0rdX"})
        assert r.status_code == 201, r.get_json()
        self.token = r.get_json()["token"]
        self.h = {"Authorization": f"Bearer {self.token}"}

    def get(self, url, **kw): return self.c.get(url, headers=self.h, **kw)
    def post(self, url, json=None, **kw): return self.c.post(url, headers=self.h, json=json if json is not None else {}, **kw)
    def delete(self, url): return self.c.delete(url, headers=self.h)

    def start_domain(self, slug="working-capital-treasury", blocks=None):
        assert self.post("/api/apps/select", {"app": "finance"}).status_code == 200
        assert self.post(f"/api/domains/{slug}/select").status_code == 200
        r = self.post(f"/api/domains/{slug}/blocks", {"blocks": blocks or ["core", "shared"]})
        assert r.status_code == 200, r.get_json()
        return r.get_json()

    def answer_until(self, slug, answers):
        """Feed answers (list of str) to successive open questions."""
        r = self.post(f"/api/domains/{slug}/conversation")
        assert r.status_code in (200, 201), r.get_json()
        snap = r.get_json()
        for a in answers:
            assert snap["open_question"], "no open question"
            r = self.post(f"/api/domains/{slug}/conversation/answer", {"message": a})
            assert r.status_code == 200, r.get_json()
            snap = r.get_json()
        return snap

    def kg1(self, slug="working-capital-treasury"):
        """Reach KG1 via the chatbot: rename + new attribute + experience node, then approve everything."""
        self.start_domain(slug)
        r = self.post(f"/api/domains/{slug}/conversation")
        snap = r.get_json()
        guard = 0
        while snap["open_question"] and guard < 30:
            guard += 1
            topic = snap["open_question"]["topic"]
            msg = {"naming": "We call it Receivables Control Desk", "attributes": "creditLimit, region code",
                   "relationships": "same", "rules": "same", "missing": "Distributor Rebate Ledger",
                   "experience": "Large distributors pay late at quarter end so we keep an extra reserve."}[topic]
            snap = self.post(f"/api/domains/{slug}/conversation/answer", {"message": msg}).get_json()
            if topic == "experience":
                break
        prop = self.post(f"/api/domains/{slug}/conversation/propose").get_json()["proposal"]
        assert prop, "expected a proposal"
        d = self.post(f"/api/domains/{slug}/proposals/{prop['id']}/decision", {"mode": "approve_all"})
        assert d.status_code == 200, d.get_json()
        return d.get_json()


@pytest.fixture()
def api(client):
    return Api(client)


@pytest.fixture()
def api2(client):
    return Api(client)


# ----------------------------------------------------------------- tiny file builders
def make_pdf(text="Debtor Position report arDso ageingBuckets"):
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
            b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream", b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offs = bytearray(b"%PDF-1.4\n"), []
    for i, o in enumerate(objs, 1):
        offs.append(len(out)); out += b"%d 0 obj\n" % i + o + b"\nendobj\n"
    x = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    for o in offs: out += b"%010d 00000 n \n" % o
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF" % (len(objs) + 1, x)
    return bytes(out)


def make_docx(paragraphs=("Working capital review",), table=None):
    import docx
    d = docx.Document()
    for p in paragraphs: d.add_paragraph(p)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for i, row in enumerate(table):
            for j, v in enumerate(row): t.cell(i, j).text = v
    b = io.BytesIO(); d.save(b); return b.getvalue()


def make_xlsx(rows):
    import pandas as pd
    b = io.BytesIO(); pd.DataFrame(rows).to_excel(b, index=False, engine="openpyxl"); return b.getvalue()
