"""File validation + extraction. Pure functions (no DB). Uploaded SQL is only ever READ AS TEXT — never executed."""
import io
import re
import zipfile
from pathlib import Path

ALLOWED_EXT = {"csv", "xlsx", "pdf", "txt", "docx", "sql"}
PREVIEW_ROWS = 10
CELL_LIMIT = 80
TEXT_LIMIT = 1_000_000


class FileRejected(Exception):
    """Content does not match the declared type / is unsafe."""


def sniff(path: Path, ext: str):
    """Validate that the bytes actually look like the declared extension."""
    with open(path, "rb") as f:
        head = f.read(8192)
    if not head:
        raise FileRejected("The file is empty.")
    if ext == "pdf":
        if not head.lstrip().startswith(b"%PDF-"):
            raise FileRejected("This is not a valid PDF (missing %PDF header).")
    elif ext in ("docx", "xlsx"):
        if not head.startswith(b"PK\x03\x04"):
            raise FileRejected(f"This is not a valid {ext.upper()} file (not an Office Open XML package).")
        try:
            with zipfile.ZipFile(path) as z:
                names = z.namelist()
        except zipfile.BadZipFile:
            raise FileRejected(f"This {ext.upper()} file is corrupted.")
        need = "word/document.xml" if ext == "docx" else "xl/workbook.xml"
        if need not in names:
            raise FileRejected(f"This ZIP is not a real {ext.upper()} document.")
    else:  # csv, txt, sql
        if b"\x00" in head:
            raise FileRejected(f"This does not look like a text-based .{ext} file (binary content found).")
        if head.startswith((b"PK\x03\x04", b"%PDF-")):
            raise FileRejected(f"The content is not plain text, so it cannot be a .{ext} file.")


def read_text(path: Path, limit=TEXT_LIMIT) -> str:
    raw = path.read_bytes()[: limit * 2]
    for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            txt = raw.decode(enc)
            if enc == "utf-16" and "\x00" in txt:
                continue
            return txt[:limit]
        except (UnicodeDecodeError, UnicodeError):
            continue
    return raw.decode("latin-1", errors="replace")[:limit]


def _cell(v):
    s = "" if v is None else str(v)
    return s if len(s) <= CELL_LIMIT else s[:CELL_LIMIT] + "…"


def _infer(series):
    import pandas as pd
    s = series[series.astype(str).str.strip() != ""].head(2000)
    if s.empty:
        return "empty"
    if pd.to_numeric(s.astype(str).str.replace(",", "", regex=False), errors="coerce").notna().mean() > 0.9:
        return "number"
    try:
        if pd.to_datetime(s, errors="coerce", format="mixed").notna().mean() > 0.9:
            return "date"
    except (TypeError, ValueError):
        pass
    return "text"


def _profile_frame(df):
    cols = []
    for c in df.columns:
        sample = [x for x in df[c].fillna("").astype(str).head(2000).tolist() if x.strip()]
        cols.append({"name": str(c), "type": _infer(df[c]), "non_empty_in_sample": len(sample),
                     "examples": [_cell(x) for x in list(dict.fromkeys(sample))[:3]]})
    return cols


def detect_delimiter(text: str) -> str:
    """Pick the delimiter that gives the most consistent column count over the first lines (default ',')."""
    import csv
    lines = [l for l in text.splitlines()[:30] if l.strip()]
    best, best_score = ",", 0
    for cand in (",", ";", "\t", "|"):
        counts = [len(r) for r in csv.reader(lines, delimiter=cand)]
        if not counts:
            continue
        modal = max(set(counts), key=counts.count)
        if modal > 1 and counts.count(modal) * modal > best_score:
            best, best_score = cand, counts.count(modal) * modal
    return best


def parse_csv(path: Path):
    import pandas as pd
    text = read_text(path, TEXT_LIMIT * 20)
    # read everything as strings so nothing is silently coerced
    df = pd.read_csv(io.StringIO(text), sep=detect_delimiter(text), dtype=str, keep_default_na=False, on_bad_lines="skip", nrows=200_000)
    df = df.fillna("")
    if df.shape[1] == 0:
        raise FileRejected("No columns could be read from this CSV.")
    cols = _profile_frame(df)
    return {"metadata": {"kind": "table", "rows_read": int(len(df)), "columns": len(df.columns), "column_profile": cols,
                         "row_cap_applied": len(df) >= 200_000},
            "preview": {"type": "table", "columns": [c["name"] for c in cols],
                        "rows": [[_cell(v) for v in r] for r in df.head(PREVIEW_ROWS).values.tolist()]},
            "terms": [c["name"] for c in cols], "text": None}


def parse_xlsx(path: Path):
    import pandas as pd
    try:
        sheets = pd.read_excel(path, sheet_name=None, dtype=str, keep_default_na=False, engine="openpyxl", nrows=200_000)
    except Exception as e:  # openpyxl raises many types for corrupt files
        raise FileRejected(f"The workbook could not be read: {type(e).__name__}")
    if not sheets:
        raise FileRejected("The workbook has no sheets.")
    meta_sheets, terms, first = [], [], None
    for name, df in sheets.items():
        df = df.loc[:, [c for c in df.columns if not str(c).startswith("Unnamed:")]] if len(df.columns) else df
        cols = _profile_frame(df) if len(df.columns) else []
        meta_sheets.append({"name": str(name), "rows_read": int(len(df)), "columns": len(df.columns), "column_profile": cols})
        terms += [c["name"] for c in cols]
        if first is None and cols:
            first = {"type": "table", "sheet": str(name), "columns": [c["name"] for c in cols],
                     "rows": [[_cell(v) for v in r] for r in df[[c["name"] for c in cols]].head(PREVIEW_ROWS).values.tolist()]}
    return {"metadata": {"kind": "workbook", "sheets": meta_sheets, "sheet_count": len(meta_sheets)},
            "preview": first or {"type": "text", "text": "No tabular data found in this workbook."},
            "terms": list(dict.fromkeys(terms)), "text": None}


def parse_pdf(path: Path):
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError
    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            raise FileRejected("The PDF is password-protected, so it cannot be read.")
        pages = len(reader.pages)
        chunks = []
        for p in reader.pages[:60]:
            chunks.append(p.extract_text() or "")
    except PdfReadError as e:
        raise FileRejected(f"The PDF could not be read: {e}")
    text = "\n".join(chunks).strip()
    meta = {"kind": "document", "pages": pages, "pages_read": min(pages, 60), "text_chars": len(text)}
    if not text:
        meta["warning"] = "No extractable text found (the PDF may be a scanned image). It cannot be mapped to the graph."
    return {"metadata": meta, "preview": {"type": "text", "text": text[:1500] or "(no extractable text)"}, "terms": [], "text": text[:TEXT_LIMIT]}


def parse_docx(path: Path):
    import docx
    try:
        d = docx.Document(str(path))
    except Exception as e:
        raise FileRejected(f"The Word document could not be read: {type(e).__name__}")
    paras = [p.text for p in d.paragraphs if p.text.strip()]
    table_terms, table_text = [], []
    for t in d.tables:
        for ri, row in enumerate(t.rows):
            cells = [c.text.strip() for c in row.cells]
            table_text.append(" | ".join(cells))
            if ri == 0:
                table_terms += [c for c in cells if c]
    text = "\n".join(paras + table_text)
    meta = {"kind": "document", "paragraphs": len(paras), "tables": len(d.tables), "text_chars": len(text)}
    if not text.strip():
        meta["warning"] = "The document contains no text."
    return {"metadata": meta, "preview": {"type": "text", "text": text[:1500] or "(empty document)"},
            "terms": list(dict.fromkeys(table_terms)), "text": text[:TEXT_LIMIT]}


def parse_txt(path: Path):
    text = read_text(path)
    if not text.strip():
        raise FileRejected("The text file is empty.")
    return {"metadata": {"kind": "document", "text_chars": len(text), "lines": text.count("\n") + 1},
            "preview": {"type": "text", "text": text[:1500]}, "terms": [], "text": text}


_CREATE = re.compile(r"create\s+(?:or\s+replace\s+)?(?:temp(?:orary)?\s+)?table\s+(?:if\s+not\s+exists\s+)?[`\"\[]?(?:\w+[`\"\]]?\.[`\"\[]?)?(\w+)[`\"\]]?\s*\((.*)\)", re.I | re.S)
_INSERT = re.compile(r"insert\s+(?:ignore\s+)?into\s+[`\"\[]?(?:\w+[`\"\]]?\.[`\"\[]?)?(\w+)[`\"\]]?\s*\(([^)]*)\)", re.I)
_SELECT = re.compile(r"select\s+(.*?)\s+from\s+[`\"\[]?(?:\w+[`\"\]]?\.[`\"\[]?)?(\w+)", re.I | re.S)
_CONSTRAINT = ("primary", "foreign", "unique", "key", "constraint", "index", "check", "fulltext")


def split_sql(text):
    """Split on ';' outside quotes/comments (string handling only; the SQL is never run)."""
    out, buf, q, i = [], [], None, 0
    while i < len(text):
        c = text[i]
        if q:
            buf.append(c)
            if c == q and text[i - 1] != "\\":
                q = None
        elif c in ("'", '"', "`"):
            q = c; buf.append(c)
        elif text.startswith("--", i):
            j = text.find("\n", i); i = len(text) if j < 0 else j; continue
        elif text.startswith("/*", i):
            j = text.find("*/", i + 2); i = len(text) if j < 0 else j + 2; continue
        elif c == ";":
            s = "".join(buf).strip()
            if s: out.append(s)
            buf = []
        else:
            buf.append(c)
        i += 1
    s = "".join(buf).strip()
    if s:
        out.append(s)
    return out


def parse_sql(path: Path):
    text = read_text(path)
    stmts = split_sql(text)
    if not stmts:
        raise FileRejected("No SQL statements were found in this file.")
    tables, counts = {}, {}
    for s in stmts:
        kind = (s.split(None, 1)[0] or "").upper()
        counts[kind] = counts.get(kind, 0) + 1
        m = _CREATE.search(s)
        if m:
            cols = []
            depth, cur = 0, ""
            for ch in m.group(2) + ",":
                if ch == "(": depth += 1
                if ch == ")": depth -= 1
                if ch == "," and depth == 0:
                    first = cur.strip().split()
                    if first and first[0].strip("`\"[]").lower() not in _CONSTRAINT:
                        cols.append(first[0].strip("`\"[]"))
                    cur = ""
                else:
                    cur += ch
            tables.setdefault(m.group(1), [])
            tables[m.group(1)] += [c for c in cols if c not in tables[m.group(1)]]
            continue
        m = _INSERT.search(s)
        if m:
            cols = [c.strip().strip("`\"[]") for c in m.group(2).split(",") if c.strip()]
            t = tables.setdefault(m.group(1), [])
            t += [c for c in cols if c not in t]
            continue
        m = _SELECT.match(s) if kind == "SELECT" else None
        if m:
            cols = [c.strip().split(" as ")[-1].split(".")[-1].strip("`\"[] ") for c in m.group(1).split(",") if c.strip() and c.strip() != "*"]
            t = tables.setdefault(m.group(2), [])
            t += [c for c in cols if re.fullmatch(r"\w+", c) and c not in t]
    terms = list(dict.fromkeys([t for t in tables] + [c for cols in tables.values() for c in cols]))
    meta = {"kind": "sql", "statements": len(stmts), "statement_types": counts, "tables": {t: c[:60] for t, c in list(tables.items())[:60]},
            "executed": False, "note": "SQL files are parsed as text only and are never executed."}
    if not tables:
        meta["warning"] = "No table or column definitions could be detected in this SQL."
    return {"metadata": meta, "preview": {"type": "sql", "tables": meta["tables"], "text": text[:800]}, "terms": terms, "text": None}


PARSERS = {"csv": parse_csv, "xlsx": parse_xlsx, "pdf": parse_pdf, "docx": parse_docx, "txt": parse_txt, "sql": parse_sql}
